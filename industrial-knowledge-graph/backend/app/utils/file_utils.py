import os
import re
import logging
from pathlib import Path
from app.config import settings

logger = logging.getLogger(__name__)

ALLOWED_EXTENSIONS = {'.pdf', '.docx', '.txt', '.csv', '.xlsx'}

def validate_file(filename: str, content_type: str = None) -> bool:
    """Check that the file has an allowed extension."""
    ext = os.path.splitext(filename)[1].lower()
    return ext in ALLOWED_EXTENSIONS

def safe_filename(filename: str) -> str:
    """Sanitise a filename to only contain safe characters."""
    return re.sub(r'[^a-zA-Z0-9_.\-]', '_', filename)

def get_upload_path(filename: str) -> Path:
    """Return a unique upload path for the given filename."""
    import uuid
    stem = Path(safe_filename(filename)).stem
    suffix = Path(filename).suffix.lower()
    unique_name = f"{stem}_{uuid.uuid4().hex[:8]}{suffix}"
    return Path(settings.UPLOAD_DIR) / unique_name

# ── Page Extractors ────────────────────────────────────────────────────────────

def extract_pages_from_pdf(file_path: str) -> list:
    pages = []
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(file_path)
        for idx, page in enumerate(doc):
            p_text = clean_text(page.get_text())
            if p_text:
                pages.append({"page_number": idx + 1, "text": p_text})
        doc.close()
        logger.info(f"Extracted {len(pages)} pages from PDF via PyMuPDF: {file_path}")
    except Exception as e:
        try:
            import PyPDF2
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                for idx, page in enumerate(reader.pages):
                    p_text = clean_text(page.extract_text() or "")
                    if p_text:
                        pages.append({"page_number": idx + 1, "text": p_text})
            logger.info(f"Extracted {len(pages)} pages from PDF via PyPDF2: {file_path}")
        except Exception as ex:
            logger.error(f"PDF page extraction failed: {ex}")
    return pages

def extract_pages_from_docx(file_path: str) -> list:
    pages = []
    try:
        import docx
        doc = docx.Document(file_path)
        full_text_blocks = []
        for para in doc.paragraphs:
            if para.text.strip():
                full_text_blocks.append(para.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    full_text_blocks.append(row_text)

        # Virtual page split: ~350 words per page
        words = ("\n".join(full_text_blocks)).split()
        page_size = 350
        page_num = 1
        for i in range(0, len(words), page_size):
            p_words = words[i:i + page_size]
            p_text = clean_text(" ".join(p_words))
            if p_text:
                pages.append({"page_number": page_num, "text": p_text})
                page_num += 1
    except Exception as e:
        logger.error(f"DOCX page extraction failed: {e}")
    return pages

def extract_pages_from_txt(file_path: str) -> list:
    pages = []
    try:
        with open(file_path, 'r', encoding='utf-8', errors='replace') as f:
            content = f.read()
        words = content.split()
        page_size = 350
        page_num = 1
        for i in range(0, max(1, len(words)), page_size):
            p_words = words[i:i + page_size]
            p_text = clean_text(" ".join(p_words))
            if p_text:
                pages.append({"page_number": page_num, "text": p_text})
                page_num += 1
    except Exception as e:
        logger.error(f"TXT page extraction failed: {e}")
    return pages

def extract_pages_from_csv(file_path: str) -> list:
    pages = []
    try:
        import pandas as pd
        df = pd.read_csv(file_path, encoding='utf-8', errors='replace')
        rows_per_page = 25
        page_num = 1
        for i in range(0, len(df), rows_per_page):
            chunk_df = df.iloc[i:i+rows_per_page]
            p_text = f"CSV Section (Rows {i+1} to {i+len(chunk_df)}):\n"
            for _, row in chunk_df.iterrows():
                row_parts = [f"{col}: {val}" for col, val in row.items() if str(val).strip() and str(val) != 'nan']
                if row_parts:
                    p_text += " | ".join(row_parts) + "\n"
            pages.append({"page_number": page_num, "text": clean_text(p_text)})
            page_num += 1
    except Exception as e:
        logger.error(f"CSV page extraction failed: {e}")
    return pages

def extract_pages_from_xlsx(file_path: str) -> list:
    pages = []
    try:
        import pandas as pd
        xl = pd.ExcelFile(file_path)
        page_num = 1
        for sheet_name in xl.sheet_names:
            df = xl.parse(sheet_name)
            p_text = f"=== Sheet: {sheet_name} ===\n"
            p_text += f"Columns: {', '.join(str(c) for c in df.columns.tolist())}\n\n"
            for _, row in df.iterrows():
                row_parts = [f"{col}: {val}" for col, val in row.items() if str(val).strip() and str(val) != 'nan']
                if row_parts:
                    p_text += " | ".join(row_parts) + "\n"
            pages.append({"page_number": page_num, "text": clean_text(p_text)})
            page_num += 1
    except Exception as e:
        logger.error(f"XLSX page extraction failed: {e}")
    return pages

def extract_pages(file_path: str, file_type: str = None) -> list:
    """Dispatch to correct extractor and return list of {'page_number': N, 'text': '...'}"""
    ext = os.path.splitext(file_path)[1].lower()
    extractors = {
        '.pdf':  extract_pages_from_pdf,
        '.docx': extract_pages_from_docx,
        '.txt':  extract_pages_from_txt,
        '.csv':  extract_pages_from_csv,
        '.xlsx': extract_pages_from_xlsx,
    }
    extractor = extractors.get(ext)
    if not extractor:
        return []
    return extractor(file_path)

def extract_text(file_path: str, file_type: str = None) -> str:
    pages = extract_pages(file_path, file_type)
    return "\n\n".join([f"--- Page {p['page_number']} ---\n" + p['text'] for p in pages])

def split_into_chunks(text: str, chunk_size: int = 800, overlap: int = 150) -> list:
    words = text.split()
    chunks = []
    if not words:
        return chunks
    i = 0
    while i < len(words):
        chunk_words = words[i:i + chunk_size]
        chunks.append(" ".join(chunk_words))
        i += chunk_size - overlap
        if i + overlap >= len(words) and i < len(words):
            chunks.append(" ".join(words[i:]))
            break
    return chunks

def clean_text(text: str) -> str:
    if not text: return ""
    text = text.replace('\x00', '')
    text = re.sub(r'\r\n|\r', '\n', text)
    text = re.sub(r'\n{3,}', '\n\n', text)
    text = re.sub(r'[ \t]{2,}', ' ', text)
    return text.strip()
